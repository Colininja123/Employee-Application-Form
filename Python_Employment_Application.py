import tkinter as tk
from tkinter import messagebox
import re
import os
import openpyxl

root = tk.Tk()
root.title("Employment Application - Page 1")

# ---------- Window Size + Center ----------
window_width = 950
window_height = 620

screen_width = root.winfo_screenwidth()
screen_height = root.winfo_screenheight()

x = int((screen_width / 2) - (window_width / 2))
y = int((screen_height / 2) - (window_height / 2))

root.geometry(f"{window_width}x{window_height}+{x}+{y}")

# ---------- Variables Page 1 ----------
first_name = tk.StringVar()
last_name = tk.StringVar()
street = tk.StringVar()
city = tk.StringVar()
state = tk.StringVar()
zip_code = tk.StringVar()
phone = tk.StringVar()
email = tk.StringVar()
date_available = tk.StringVar()
position = tk.StringVar()

# ---------- Variables Page 2 ----------
from_date1 = tk.StringVar()
to_date1 = tk.StringVar()
employer_name1 = tk.StringVar()
reason1 = tk.StringVar()

from_date2 = tk.StringVar()
to_date2 = tk.StringVar()
employer_name2 = tk.StringVar()
reason2 = tk.StringVar()

from_date3 = tk.StringVar()
to_date3 = tk.StringVar()
employer_name3 = tk.StringVar()
reason3 = tk.StringVar()

from_date4 = tk.StringVar()
to_date4 = tk.StringVar()
employer_name4 = tk.StringVar()
reason4 = tk.StringVar()

# ---------- Variables Page 3 ----------
school_started1 = tk.StringVar()
school_ended1 = tk.StringVar()
school_name1 = tk.StringVar()
degree1 = tk.StringVar()

school_started2 = tk.StringVar()
school_ended2 = tk.StringVar()
school_name2 = tk.StringVar()
degree2 = tk.StringVar()

school_started3 = tk.StringVar()
school_ended3 = tk.StringVar()
school_name3 = tk.StringVar()
degree3 = tk.StringVar()

school_started4 = tk.StringVar()
school_ended4 = tk.StringVar()
school_name4 = tk.StringVar()
degree4 = tk.StringVar()

# ---------- Variables Page 4 ----------
ref_name1 = tk.StringVar()
ref_email1 = tk.StringVar()
ref_phone1 = tk.StringVar()

ref_name2 = tk.StringVar()
ref_email2 = tk.StringVar()
ref_phone2 = tk.StringVar()

ref_name3 = tk.StringVar()
ref_email3 = tk.StringVar()
ref_phone3 = tk.StringVar()

# ---------- Frames ----------
page1 = tk.Frame(root)
page2 = tk.Frame(root)
page3 = tk.Frame(root)
page4 = tk.Frame(root)

for frame in (page1, page2, page3, page4):
    frame.place(x=0, y=0, relwidth=1, relheight=1)

# ---------- Placeholder Function ----------
def add_placeholder(entry, text):
    entry.insert(0, text)
    entry.config(fg="gray")

    def on_focus_in(event):
        if entry.get() == text:
            entry.delete(0, tk.END)
            entry.config(fg="black")

    def on_focus_out(event):
        if entry.get() == "":
            entry.insert(0, text)
            entry.config(fg="gray")

    entry.bind("<FocusIn>", on_focus_in)
    entry.bind("<FocusOut>", on_focus_out)

# ---------- ZIP Limit ----------
def limit_zip(*args):
    value = zip_code.get()
    if len(value) > 5:
        zip_code.set(value[:5])

zip_code.trace_add("write", limit_zip)

# ---------- Validation ----------
def validate_page1():
    fields = [
        first_name.get().strip(),
        last_name.get().strip(),
        street.get().strip(),
        city.get().strip(),
        state.get().strip(),
        zip_code.get().strip(),
        phone.get().strip(),
        email.get().strip(),
        date_available.get().strip(),
        position.get().strip()
    ]

    placeholders = ["(999) 999 9999", "name@example.com", "MM/DD/YYYY"]

    if (
        "" in fields
        or "Select State" in fields
        or "Select Position" in fields
        or phone.get().strip() in placeholders
        or email.get().strip() in placeholders
        or date_available.get().strip() in placeholders
    ):
        messagebox.showerror("Error", "All fields on Page 1 are required")
        return False

    if not re.fullmatch(r"\d{5}", zip_code.get().strip()):
        messagebox.showerror("Error", "Zip Code must be 5 digits")
        return False

    if not re.fullmatch(r"\(\d{3}\) \d{3} \d{4}", phone.get().strip()):
        messagebox.showerror("Error", "Phone Number must be in this format: (999) 999 9999")
        return False

    if not re.fullmatch(r"[^@]+@[^@]+\.[^@]+", email.get().strip()):
        messagebox.showerror("Error", "Enter a valid email address like name@example.com")
        return False

    if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", date_available.get().strip()):
        messagebox.showerror("Error", "Date Available must be in MM/DD/YYYY format")
        return False

    return True

def validate_one_employer(from_date, to_date, employer_name, reason, number, required=False):
    fd = from_date.get().strip()
    td = to_date.get().strip()
    en = employer_name.get().strip()
    rs = reason.get().strip()

    empty_reason = rs == "Select Reason"
    values = [fd, td, en]
    filled_count = sum(1 for value in values if value != "")

    if required:
        if fd == "" or td == "" or en == "" or empty_reason:
            messagebox.showerror("Error", f"Employer {number} is required")
            return False
    else:
        if filled_count == 0 and empty_reason:
            return True
        if filled_count < 3 or empty_reason:
            messagebox.showerror("Error", f"Employer {number}: if one field is entered, all related fields are required")
            return False

    if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", fd):
        messagebox.showerror("Error", f"Employer {number}: From Date must be MM/DD/YYYY")
        return False

    if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", td):
        messagebox.showerror("Error", f"Employer {number}: To Date must be MM/DD/YYYY")
        return False

    return True

def validate_page2():
    if not validate_one_employer(from_date1, to_date1, employer_name1, reason1, 1, required=True):
        return False
    if not validate_one_employer(from_date2, to_date2, employer_name2, reason2, 2):
        return False
    if not validate_one_employer(from_date3, to_date3, employer_name3, reason3, 3):
        return False
    if not validate_one_employer(from_date4, to_date4, employer_name4, reason4, 4):
        return False
    return True

def validate_one_school(started, ended, school_name, degree, number, required=False):
    sd = started.get().strip()
    ed = ended.get().strip()
    sn = school_name.get().strip()
    dg = degree.get().strip()

    if sn == "School Name":
        sn = ""

    empty_degree = dg == "Select One"
    values = [sd, ed, sn]
    filled_count = sum(1 for value in values if value != "")

    if required:
        if sd == "" or ed == "" or sn == "" or empty_degree:
            messagebox.showerror("Error", f"School {number} is required")
            return False
    else:
        if filled_count == 0 and empty_degree:
            return True
        if filled_count < 3 or empty_degree:
            messagebox.showerror("Error", f"School {number}: if one field is entered, all related fields are required")
            return False

    if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", sd):
        messagebox.showerror("Error", f"School {number}: Date Started must be MM/DD/YYYY")
        return False

    if not re.fullmatch(r"\d{2}/\d{2}/\d{4}", ed):
        messagebox.showerror("Error", f"School {number}: Date Ended must be MM/DD/YYYY")
        return False

    return True

def validate_page3():
    if not validate_one_school(school_started1, school_ended1, school_name1, degree1, 1, required=True):
        return False
    if not validate_one_school(school_started2, school_ended2, school_name2, degree2, 2):
        return False
    if not validate_one_school(school_started3, school_ended3, school_name3, degree3, 3):
        return False
    if not validate_one_school(school_started4, school_ended4, school_name4, degree4, 4):
        return False
    return True

def validate_one_reference(full_name, email_addr, phone_num, number):
    fn = full_name.get().strip()
    em = email_addr.get().strip()
    ph = phone_num.get().strip()

    placeholders = ["First Last", "example@email.com", "(555) 555 5555"]

    if fn == "" or em == "" or ph == "" or fn in placeholders or em in placeholders or ph in placeholders:
        messagebox.showerror("Error", f"All fields for Reference {number} are required")
        return False

    if not re.fullmatch(r"[^@]+@[^@]+\.[^@]+", em):
        messagebox.showerror("Error", f"Reference {number}: Enter a valid email address")
        return False

    if not re.fullmatch(r"\(\d{3}\) \d{3} \d{4}", ph):
        messagebox.showerror("Error", f"Reference {number}: Phone Number must be in this format: (555) 555 5555")
        return False

    return True

def validate_page4():
    if not validate_one_reference(ref_name1, ref_email1, ref_phone1, 1):
        return False
    if not validate_one_reference(ref_name2, ref_email2, ref_phone2, 2):
        return False
    if not validate_one_reference(ref_name3, ref_email3, ref_phone3, 3):
        return False
    return True

# ---------- Navigation ----------
def go_to_page2():
    if validate_page1():
        root.title("Employment Application - Page 2")
        applicant_name_label.config(text=f"Applicant: {first_name.get().strip()} {last_name.get().strip()}")
        page2.tkraise()

def go_to_page1():
    root.title("Employment Application - Page 1")
    page1.tkraise()

def go_to_page3():
    if validate_page2():
        root.title("Employment Application - Page 3")
        education_applicant_label.config(text=f"Applicant: {first_name.get().strip()} {last_name.get().strip()}")
        page3.tkraise()

def go_back_to_page2():
    root.title("Employment Application - Page 2")
    page2.tkraise()

def go_to_page4():
    if validate_page3():
        root.title("Employment Application - Page 4")
        references_applicant_label.config(text=f"Applicant: {first_name.get().strip()} {last_name.get().strip()}")
        page4.tkraise()

def go_back_to_page3():
    root.title("Employment Application - Page 3")
    page3.tkraise()

# ---------- Excel Save ----------
def clear_all_fields():
    # Page 1
    first_name.set("")
    last_name.set("")
    street.set("")
    city.set("")
    state.set("Select State")
    zip_code.set("")
    phone.set("")
    email.set("")
    date_available.set("")
    position.set("Select Position")

    # Page 2
    from_date1.set("")
    to_date1.set("")
    employer_name1.set("")
    reason1.set("Select Reason")

    from_date2.set("")
    to_date2.set("")
    employer_name2.set("")
    reason2.set("Select Reason")

    from_date3.set("")
    to_date3.set("")
    employer_name3.set("")
    reason3.set("Select Reason")

    from_date4.set("")
    to_date4.set("")
    employer_name4.set("")
    reason4.set("Select Reason")

    # Page 3
    school_started1.set("")
    school_ended1.set("")
    school_name1.set("")
    degree1.set("Select One")

    school_started2.set("")
    school_ended2.set("")
    school_name2.set("")
    degree2.set("Select One")

    school_started3.set("")
    school_ended3.set("")
    school_name3.set("")
    degree3.set("Select One")

    school_started4.set("")
    school_ended4.set("")
    school_name4.set("")
    degree4.set("Select One")

    # Page 4
    ref_name1.set("")
    ref_email1.set("")
    ref_phone1.set("")

    ref_name2.set("")
    ref_email2.set("")
    ref_phone2.set("")

    ref_name3.set("")
    ref_email3.set("")
    ref_phone3.set("")

    # restore placeholders
    phone_entry.delete(0, tk.END)
    add_placeholder(phone_entry, "(999) 999 9999")

    email_entry.delete(0, tk.END)
    add_placeholder(email_entry, "name@example.com")

    date_entry.delete(0, tk.END)
    add_placeholder(date_entry, "MM/DD/YYYY")

    from1_entry.delete(0, tk.END)
    add_placeholder(from1_entry, "MM/DD/YYYY")
    to1_entry.delete(0, tk.END)
    add_placeholder(to1_entry, "MM/DD/YYYY")

    from2_entry.delete(0, tk.END)
    add_placeholder(from2_entry, "MM/DD/YYYY")
    to2_entry.delete(0, tk.END)
    add_placeholder(to2_entry, "MM/DD/YYYY")

    from3_entry.delete(0, tk.END)
    add_placeholder(from3_entry, "MM/DD/YYYY")
    to3_entry.delete(0, tk.END)
    add_placeholder(to3_entry, "MM/DD/YYYY")

    from4_entry.delete(0, tk.END)
    add_placeholder(from4_entry, "MM/DD/YYYY")
    to4_entry.delete(0, tk.END)
    add_placeholder(to4_entry, "MM/DD/YYYY")

    school1_start_entry.delete(0, tk.END)
    add_placeholder(school1_start_entry, "MM/DD/YYYY")
    school1_end_entry.delete(0, tk.END)
    add_placeholder(school1_end_entry, "MM/DD/YYYY")
    school1_name_entry.delete(0, tk.END)
    add_placeholder(school1_name_entry, "School Name")

    school2_start_entry.delete(0, tk.END)
    add_placeholder(school2_start_entry, "MM/DD/YYYY")
    school2_end_entry.delete(0, tk.END)
    add_placeholder(school2_end_entry, "MM/DD/YYYY")
    school2_name_entry.delete(0, tk.END)
    add_placeholder(school2_name_entry, "School Name")

    school3_start_entry.delete(0, tk.END)
    add_placeholder(school3_start_entry, "MM/DD/YYYY")
    school3_end_entry.delete(0, tk.END)
    add_placeholder(school3_end_entry, "MM/DD/YYYY")
    school3_name_entry.delete(0, tk.END)
    add_placeholder(school3_name_entry, "School Name")

    school4_start_entry.delete(0, tk.END)
    add_placeholder(school4_start_entry, "MM/DD/YYYY")
    school4_end_entry.delete(0, tk.END)
    add_placeholder(school4_end_entry, "MM/DD/YYYY")
    school4_name_entry.delete(0, tk.END)
    add_placeholder(school4_name_entry, "School Name")

    ref1_name_entry.delete(0, tk.END)
    add_placeholder(ref1_name_entry, "First Last")
    ref1_email_entry.delete(0, tk.END)
    add_placeholder(ref1_email_entry, "example@email.com")
    ref1_phone_entry.delete(0, tk.END)
    add_placeholder(ref1_phone_entry, "(555) 555 5555")

    ref2_name_entry.delete(0, tk.END)
    add_placeholder(ref2_name_entry, "First Last")
    ref2_email_entry.delete(0, tk.END)
    add_placeholder(ref2_email_entry, "example@email.com")
    ref2_phone_entry.delete(0, tk.END)
    add_placeholder(ref2_phone_entry, "(555) 555 5555")

    ref3_name_entry.delete(0, tk.END)
    add_placeholder(ref3_name_entry, "First Last")
    ref3_email_entry.delete(0, tk.END)
    add_placeholder(ref3_email_entry, "example@email.com")
    ref3_phone_entry.delete(0, tk.END)
    add_placeholder(ref3_phone_entry, "(555) 555 5555")

def submit_application():
    if not validate_page4():
        return

    filename = "employment_application.xlsx"

    if not os.path.exists(filename):
        wb = openpyxl.Workbook()

        applicants_sheet = wb.active
        applicants_sheet.title = "Applicants"
        applicants_sheet.append([
            "First Name", "Last Name", "Street Address", "City", "State",
            "Zip Code", "Phone Number", "Email Address", "Date Available",
            "Desired Position"
        ])

        work_history_sheet = wb.create_sheet("Work History")
        work_history_sheet.append([
            "Applicant Name", "Employer", "From Date", "To Date",
            "Employer Name", "Reason For Leaving"
        ])

        education_sheet = wb.create_sheet("Education")
        education_sheet.append([
            "Applicant Name", "School", "Date Started", "Date Ended",
            "School Name", "Degree Awarded"
        ])

        references_sheet = wb.create_sheet("References")
        references_sheet.append([
            "Applicant Name", "Reference", "Full Name", "Email Address", "Phone Number"
        ])

        wb.save(filename)

    wb = openpyxl.load_workbook(filename)

    if "Applicants" not in wb.sheetnames:
        applicants_sheet = wb.create_sheet("Applicants")
        applicants_sheet.append([
            "First Name", "Last Name", "Street Address", "City", "State",
            "Zip Code", "Phone Number", "Email Address", "Date Available",
            "Desired Position"
        ])
    else:
        applicants_sheet = wb["Applicants"]

    if "Work History" not in wb.sheetnames:
        work_history_sheet = wb.create_sheet("Work History")
        work_history_sheet.append([
            "Applicant Name", "Employer", "From Date", "To Date",
            "Employer Name", "Reason For Leaving"
        ])
    else:
        work_history_sheet = wb["Work History"]

    if "Education" not in wb.sheetnames:
        education_sheet = wb.create_sheet("Education")
        education_sheet.append([
            "Applicant Name", "School", "Date Started", "Date Ended",
            "School Name", "Degree Awarded"
        ])
    else:
        education_sheet = wb["Education"]

    if "References" not in wb.sheetnames:
        references_sheet = wb.create_sheet("References")
        references_sheet.append([
            "Applicant Name", "Reference", "Full Name", "Email Address", "Phone Number"
        ])
    else:
        references_sheet = wb["References"]

    applicants_sheet.append([
        first_name.get().strip(),
        last_name.get().strip(),
        street.get().strip(),
        city.get().strip(),
        state.get().strip(),
        zip_code.get().strip(),
        phone.get().strip(),
        email.get().strip(),
        date_available.get().strip(),
        position.get().strip()
    ])

    applicant_name = f"{first_name.get().strip()} {last_name.get().strip()}"

    work_rows = [
        ("Employer 1", from_date1.get().strip(), to_date1.get().strip(), employer_name1.get().strip(), reason1.get().strip()),
        ("Employer 2", from_date2.get().strip(), to_date2.get().strip(), employer_name2.get().strip(), reason2.get().strip()),
        ("Employer 3", from_date3.get().strip(), to_date3.get().strip(), employer_name3.get().strip(), reason3.get().strip()),
        ("Employer 4", from_date4.get().strip(), to_date4.get().strip(), employer_name4.get().strip(), reason4.get().strip())
    ]

    for employer_label, fd, td, en, rs in work_rows:
        if fd != "" or td != "" or en != "" or rs != "Select Reason":
            work_history_sheet.append([
                applicant_name,
                employer_label,
                fd,
                td,
                en,
                rs
            ])

    school_rows = [
        ("School 1", school_started1.get().strip(), school_ended1.get().strip(), school_name1.get().strip(), degree1.get().strip()),
        ("School 2", school_started2.get().strip(), school_ended2.get().strip(), school_name2.get().strip(), degree2.get().strip()),
        ("School 3", school_started3.get().strip(), school_ended3.get().strip(), school_name3.get().strip(), degree3.get().strip()),
        ("School 4", school_started4.get().strip(), school_ended4.get().strip(), school_name4.get().strip(), degree4.get().strip())
    ]

    for school_label, sd, ed, sn, dg in school_rows:
        if sn == "School Name":
            sn = ""
        if sd != "" or ed != "" or sn != "" or dg != "Select One":
            education_sheet.append([
                applicant_name,
                school_label,
                sd,
                ed,
                sn,
                dg
            ])

    reference_rows = [
        ("Reference 1", ref_name1.get().strip(), ref_email1.get().strip(), ref_phone1.get().strip()),
        ("Reference 2", ref_name2.get().strip(), ref_email2.get().strip(), ref_phone2.get().strip()),
        ("Reference 3", ref_name3.get().strip(), ref_email3.get().strip(), ref_phone3.get().strip())
    ]

    for ref_label, fn, em, ph in reference_rows:
        references_sheet.append([
            applicant_name,
            ref_label,
            fn,
            em,
            ph
        ])

    wb.save(filename)
    messagebox.showinfo("Success", "Application Submitted Successfully!")

    clear_all_fields()
    root.title("Employment Application - Page 1")
    page1.tkraise()

# =========================
# PAGE 1
# =========================
for i in range(4):
    page1.grid_columnconfigure(i, weight=1)

tk.Label(
    page1,
    text="Employment Application - General Information",
    font=("Times New Roman", 18, "bold")
).grid(row=0, column=0, columnspan=4, pady=20)

tk.Label(page1, text="First Name:").grid(row=1, column=0, sticky="e", padx=10, pady=5)
tk.Entry(page1, textvariable=first_name, width=30).grid(row=1, column=1, sticky="w", padx=(0, 10), pady=5)

tk.Label(page1, text="Street Address:").grid(row=2, column=0, sticky="e", padx=10, pady=5)
tk.Entry(page1, textvariable=street, width=30).grid(row=2, column=1, sticky="w", padx=(0, 10), pady=5)

tk.Label(page1, text="State:").grid(row=3, column=0, sticky="e", padx=10, pady=5)
states = [
    "Select State","AL","AK","AZ","AR","CA","CO","CT","DE","FL","GA",
    "HI","ID","IL","IN","IA","KS","KY","LA","ME","MD","MA","MI","MN",
    "MS","MO","MT","NE","NV","NH","NJ","NM","NY","NC","ND","OH","OK",
    "OR","PA","RI","SC","SD","TN","TX","UT","VT","VA","WA","WV","WI","WY"
]
state.set(states[0])
state_menu = tk.OptionMenu(page1, state, *states)
state_menu.config(width=27)
state_menu.grid(row=3, column=1, sticky="w", padx=10, pady=5)

tk.Label(page1, text="Phone Number:").grid(row=4, column=0, sticky="e", padx=10, pady=5)
phone_entry = tk.Entry(page1, textvariable=phone, width=30)
phone_entry.grid(row=4, column=1, sticky="w", padx=(0, 10), pady=5)
add_placeholder(phone_entry, "(999) 999 9999")

tk.Label(page1, text="Date Available:").grid(row=5, column=0, sticky="e", padx=10, pady=5)
date_entry = tk.Entry(page1, textvariable=date_available, width=30)
date_entry.grid(row=5, column=1, sticky="w", padx=(0, 10), pady=5)
add_placeholder(date_entry, "MM/DD/YYYY")

tk.Label(page1, text="Desired Position:").grid(row=6, column=0, sticky="e", padx=10, pady=5)
positions = [
    "Select Position",
    "Office Assistant",
    "Customer Service Representative",
    "Sales Associate",
    "IT Support Specialist",
    "Manager",
    "Other"
]
position.set(positions[0])
position_menu = tk.OptionMenu(page1, position, *positions)
position_menu.config(width=35)
position_menu.grid(row=6, column=1, sticky="w", padx=10, pady=5)

tk.Label(page1, text="Last Name:").grid(row=1, column=2, sticky="e", padx=10, pady=5)
tk.Entry(page1, textvariable=last_name, width=30).grid(row=1, column=3, padx=10, pady=5)

tk.Label(page1, text="City:").grid(row=2, column=2, sticky="e", padx=10, pady=5)
tk.Entry(page1, textvariable=city, width=30).grid(row=2, column=3, padx=10, pady=5)

tk.Label(page1, text="ZIP Code:").grid(row=3, column=2, sticky="e", padx=10, pady=5)
tk.Entry(page1, textvariable=zip_code, width=30).grid(row=3, column=3, padx=10, pady=5)

tk.Label(page1, text="Email Address:").grid(row=4, column=2, sticky="e", padx=10, pady=5)
email_entry = tk.Entry(page1, textvariable=email, width=30)
email_entry.grid(row=4, column=3, padx=10, pady=5)
add_placeholder(email_entry, "name@example.com")

next_button = tk.Button(
    page1,
    text="Next",
    command=go_to_page2,
    width=15,
    font=("Arial", 11, "bold")
)
next_button.grid(row=8, column=3, sticky="e", padx=40, pady=40)

# =========================
# PAGE 2
# =========================
for i in range(5):
    page2.grid_columnconfigure(i, weight=1)

tk.Label(
    page2,
    text="Employment Application - Work History",
    font=("Times New Roman", 18, "bold")
).grid(row=0, column=0, columnspan=5, pady=(20, 10))

applicant_name_label = tk.Label(
    page2,
    text="Applicant:",
    font=("Arial", 12, "bold")
)
applicant_name_label.grid(row=1, column=0, columnspan=5, pady=(0, 20))

tk.Label(page2, text="Employer", font=("Arial", 11, "bold")).grid(row=2, column=0, padx=10, pady=5)
tk.Label(page2, text="From Date", font=("Arial", 11, "bold")).grid(row=2, column=1, padx=10, pady=5)
tk.Label(page2, text="To Date", font=("Arial", 11, "bold")).grid(row=2, column=2, padx=10, pady=5)
tk.Label(page2, text="Employer Name", font=("Arial", 11, "bold")).grid(row=2, column=3, padx=10, pady=5)
tk.Label(page2, text="Reason for Leaving", font=("Arial", 11, "bold")).grid(row=2, column=4, padx=10, pady=5)

reason_options = ["Select Reason", "Advancement Opportunity", "Relocation", "Terminated"]
reason1.set("Select Reason")
reason2.set("Select Reason")
reason3.set("Select Reason")
reason4.set("Select Reason")

tk.Label(page2, text="Employer 1:").grid(row=3, column=0, padx=(60, 10), pady=10)
from1_entry = tk.Entry(page2, textvariable=from_date1, width=12)
from1_entry.grid(row=3, column=1, padx=10, pady=10)
to1_entry = tk.Entry(page2, textvariable=to_date1, width=12)
to1_entry.grid(row=3, column=2, padx=10, pady=10)
tk.Entry(page2, textvariable=employer_name1, width=24).grid(row=3, column=3, padx=10, pady=10)
tk.OptionMenu(page2, reason1, *reason_options).grid(row=3, column=4, padx=10, pady=10)

tk.Label(page2, text="Employer 2:").grid(row=4, column=0, padx=(60, 10), pady=10)
from2_entry = tk.Entry(page2, textvariable=from_date2, width=12)
from2_entry.grid(row=4, column=1, padx=10, pady=10)
to2_entry = tk.Entry(page2, textvariable=to_date2, width=12)
to2_entry.grid(row=4, column=2, padx=10, pady=10)
tk.Entry(page2, textvariable=employer_name2, width=24).grid(row=4, column=3, padx=10, pady=10)
tk.OptionMenu(page2, reason2, *reason_options).grid(row=4, column=4, padx=10, pady=10)

tk.Label(page2, text="Employer 3:").grid(row=5, column=0, padx=(60, 10), pady=10)
from3_entry = tk.Entry(page2, textvariable=from_date3, width=12)
from3_entry.grid(row=5, column=1, padx=10, pady=10)
to3_entry = tk.Entry(page2, textvariable=to_date3, width=12)
to3_entry.grid(row=5, column=2, padx=10, pady=10)
tk.Entry(page2, textvariable=employer_name3, width=24).grid(row=5, column=3, padx=10, pady=10)
tk.OptionMenu(page2, reason3, *reason_options).grid(row=5, column=4, padx=10, pady=10)

tk.Label(page2, text="Employer 4:").grid(row=6, column=0, padx=(60, 10), pady=10)
from4_entry = tk.Entry(page2, textvariable=from_date4, width=12)
from4_entry.grid(row=6, column=1, padx=10, pady=10)
to4_entry = tk.Entry(page2, textvariable=to_date4, width=12)
to4_entry.grid(row=6, column=2, padx=10, pady=10)
tk.Entry(page2, textvariable=employer_name4, width=24).grid(row=6, column=3, padx=10, pady=10)
tk.OptionMenu(page2, reason4, *reason_options).grid(row=6, column=4, padx=10, pady=10)

tk.Button(page2, text="Back", command=go_to_page1, width=15).grid(
    row=8, column=0, sticky="w", padx=40, pady=50
)

tk.Button(page2, text="Next", command=go_to_page3, width=18).grid(
    row=8, column=4, sticky="e", padx=40, pady=50
)

# =========================
# PAGE 3
# =========================
for i in range(5):
    page3.grid_columnconfigure(i, weight=1)

tk.Label(
    page3,
    text="Employment Application - Education",
    font=("Times New Roman", 18, "bold")
).grid(row=0, column=0, columnspan=5, pady=(20, 10))

education_applicant_label = tk.Label(
    page3,
    text="Applicant:",
    font=("Arial", 12, "bold")
)
education_applicant_label.grid(row=1, column=0, columnspan=5, pady=(0, 20))

tk.Label(page3, text="School", font=("Arial", 11, "bold")).grid(row=2, column=0, padx=10, pady=5)
tk.Label(page3, text="Date Started", font=("Arial", 11, "bold")).grid(row=2, column=1, padx=10, pady=5)
tk.Label(page3, text="Date Ended", font=("Arial", 11, "bold")).grid(row=2, column=2, padx=10, pady=5)
tk.Label(page3, text="School Name", font=("Arial", 11, "bold")).grid(row=2, column=3, padx=10, pady=5)
tk.Label(page3, text="Degree Awarded", font=("Arial", 11, "bold")).grid(row=2, column=4, padx=10, pady=5)

tk.Label(page3, text="*School 1 is required", fg="red", font=("Arial", 9, "italic")).grid(
    row=3, column=0, sticky="w", padx=(20, 10), pady=(0, 0)
)

degree_options = [
    "Select One",
    "HS Diploma",
    "AAS Degree",
    "Bach Degree",
    "Masters Degree",
    "Doctorate Degree"
]

degree1.set("Select One")
degree2.set("Select One")
degree3.set("Select One")
degree4.set("Select One")

# School 1
tk.Label(page3, text="School 1:").grid(row=4, column=0, padx=(60, 10), pady=10)
school1_start_entry = tk.Entry(page3, textvariable=school_started1, width=12)
school1_start_entry.grid(row=4, column=1, padx=10, pady=10)
add_placeholder(school1_start_entry, "MM/DD/YYYY")

school1_end_entry = tk.Entry(page3, textvariable=school_ended1, width=12)
school1_end_entry.grid(row=4, column=2, padx=10, pady=10)
add_placeholder(school1_end_entry, "MM/DD/YYYY")

school1_name_entry = tk.Entry(page3, textvariable=school_name1, width=22)
school1_name_entry.grid(row=4, column=3, padx=10, pady=10)
add_placeholder(school1_name_entry, "School Name")

degree_menu1 = tk.OptionMenu(page3, degree1, *degree_options)
degree_menu1.config(width=18)
degree_menu1.grid(row=4, column=4, padx=10, pady=10)

# School 2
tk.Label(page3, text="School 2:").grid(row=5, column=0, padx=(60, 10), pady=10)
school2_start_entry = tk.Entry(page3, textvariable=school_started2, width=12)
school2_start_entry.grid(row=5, column=1, padx=10, pady=10)
add_placeholder(school2_start_entry, "MM/DD/YYYY")

school2_end_entry = tk.Entry(page3, textvariable=school_ended2, width=12)
school2_end_entry.grid(row=5, column=2, padx=10, pady=10)
add_placeholder(school2_end_entry, "MM/DD/YYYY")

school2_name_entry = tk.Entry(page3, textvariable=school_name2, width=22)
school2_name_entry.grid(row=5, column=3, padx=10, pady=10)
add_placeholder(school2_name_entry, "School Name")

degree_menu2 = tk.OptionMenu(page3, degree2, *degree_options)
degree_menu2.config(width=18)
degree_menu2.grid(row=5, column=4, padx=10, pady=10)

# School 3
tk.Label(page3, text="School 3:").grid(row=6, column=0, padx=(60, 10), pady=10)
school3_start_entry = tk.Entry(page3, textvariable=school_started3, width=12)
school3_start_entry.grid(row=6, column=1, padx=10, pady=10)
add_placeholder(school3_start_entry, "MM/DD/YYYY")

school3_end_entry = tk.Entry(page3, textvariable=school_ended3, width=12)
school3_end_entry.grid(row=6, column=2, padx=10, pady=10)
add_placeholder(school3_end_entry, "MM/DD/YYYY")

school3_name_entry = tk.Entry(page3, textvariable=school_name3, width=22)
school3_name_entry.grid(row=6, column=3, padx=10, pady=10)
add_placeholder(school3_name_entry, "School Name")

degree_menu3 = tk.OptionMenu(page3, degree3, *degree_options)
degree_menu3.config(width=18)
degree_menu3.grid(row=6, column=4, padx=10, pady=10)

# School 4
tk.Label(page3, text="School 4:").grid(row=7, column=0, padx=(60, 10), pady=10)
school4_start_entry = tk.Entry(page3, textvariable=school_started4, width=12)
school4_start_entry.grid(row=7, column=1, padx=10, pady=10)
add_placeholder(school4_start_entry, "MM/DD/YYYY")

school4_end_entry = tk.Entry(page3, textvariable=school_ended4, width=12)
school4_end_entry.grid(row=7, column=2, padx=10, pady=10)
add_placeholder(school4_end_entry, "MM/DD/YYYY")

school4_name_entry = tk.Entry(page3, textvariable=school_name4, width=22)
school4_name_entry.grid(row=7, column=3, padx=10, pady=10)
add_placeholder(school4_name_entry, "School Name")

degree_menu4 = tk.OptionMenu(page3, degree4, *degree_options)
degree_menu4.config(width=18)
degree_menu4.grid(row=7, column=4, padx=10, pady=10)

tk.Button(page3, text="Back", command=go_back_to_page2, width=15).grid(
    row=9, column=0, sticky="w", padx=40, pady=50
)

tk.Button(page3, text="Next", command=go_to_page4, width=18).grid(
    row=9, column=4, sticky="e", padx=40, pady=50
)

# =========================
# PAGE 4
# =========================
for i in range(4):
    page4.grid_columnconfigure(i, weight=1)

tk.Label(
    page4,
    text="Employment Application - References",
    font=("Times New Roman", 18, "bold")
).grid(row=0, column=0, columnspan=4, pady=(20, 10))

references_applicant_label = tk.Label(
    page4,
    text="Applicant:",
    font=("Arial", 12, "bold")
)
references_applicant_label.grid(row=1, column=0, columnspan=4, pady=(0, 5))

tk.Label(
    page4,
    text="* All 3 references are required and all fields must be completed.",
    fg="red",
    font=("Arial", 9, "italic")
).grid(row=2, column=0, columnspan=4, pady=(0, 20))

tk.Label(page4, text="Reference", font=("Arial", 11, "bold")).grid(row=3, column=0, padx=10, pady=5)
tk.Label(page4, text="Full Name", font=("Arial", 11, "bold")).grid(row=3, column=1, padx=10, pady=5)
tk.Label(page4, text="Email Address", font=("Arial", 11, "bold")).grid(row=3, column=2, padx=10, pady=5)
tk.Label(page4, text="Phone Number", font=("Arial", 11, "bold")).grid(row=3, column=3, padx=10, pady=5)

tk.Label(page4, text="Reference 1:*", fg="red", font=("Arial", 10, "bold")).grid(row=4, column=0, padx=(60, 10), pady=10)
ref1_name_entry = tk.Entry(page4, textvariable=ref_name1, width=24)
ref1_name_entry.grid(row=4, column=1, padx=10, pady=10)
add_placeholder(ref1_name_entry, "First Last")

ref1_email_entry = tk.Entry(page4, textvariable=ref_email1, width=24)
ref1_email_entry.grid(row=4, column=2, padx=10, pady=10)
add_placeholder(ref1_email_entry, "example@email.com")

ref1_phone_entry = tk.Entry(page4, textvariable=ref_phone1, width=18)
ref1_phone_entry.grid(row=4, column=3, padx=10, pady=10)
add_placeholder(ref1_phone_entry, "(555) 555 5555")

tk.Label(page4, text="Reference 2:*", fg="red", font=("Arial", 10, "bold")).grid(row=5, column=0, padx=(60, 10), pady=10)
ref2_name_entry = tk.Entry(page4, textvariable=ref_name2, width=24)
ref2_name_entry.grid(row=5, column=1, padx=10, pady=10)
add_placeholder(ref2_name_entry, "First Last")

ref2_email_entry = tk.Entry(page4, textvariable=ref_email2, width=24)
ref2_email_entry.grid(row=5, column=2, padx=10, pady=10)
add_placeholder(ref2_email_entry, "example@email.com")

ref2_phone_entry = tk.Entry(page4, textvariable=ref_phone2, width=18)
ref2_phone_entry.grid(row=5, column=3, padx=10, pady=10)
add_placeholder(ref2_phone_entry, "(555) 555 5555")

tk.Label(page4, text="Reference 3:*", fg="red", font=("Arial", 10, "bold")).grid(row=6, column=0, padx=(60, 10), pady=10)
ref3_name_entry = tk.Entry(page4, textvariable=ref_name3, width=24)
ref3_name_entry.grid(row=6, column=1, padx=10, pady=10)
add_placeholder(ref3_name_entry, "First Last")

ref3_email_entry = tk.Entry(page4, textvariable=ref_email3, width=24)
ref3_email_entry.grid(row=6, column=2, padx=10, pady=10)
add_placeholder(ref3_email_entry, "example@email.com")

ref3_phone_entry = tk.Entry(page4, textvariable=ref_phone3, width=18)
ref3_phone_entry.grid(row=6, column=3, padx=10, pady=10)
add_placeholder(ref3_phone_entry, "(555) 555 5555")

tk.Button(page4, text="Back", command=go_back_to_page3, width=15).grid(
    row=8, column=0, sticky="w", padx=40, pady=50
)

tk.Button(page4, text="Submit", command=submit_application, width=18).grid(
    row=8, column=3, sticky="e", padx=40, pady=50
)

page1.tkraise()
root.mainloop()